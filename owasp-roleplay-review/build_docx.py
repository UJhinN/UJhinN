"""Build review.docx (A4, one page) from review.md.

Usage: python3 build_docx.py [font] [size_pt]
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

FONT = sys.argv[1] if len(sys.argv) > 1 else "TH Sarabun New"
SIZE = float(sys.argv[2]) if len(sys.argv) > 2 else 14
HERE = Path(__file__).parent

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin = sec.bottom_margin = Cm(1.5)

style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(SIZE)
rpr = style.element.get_or_add_rPr()
rfonts = rpr.find(qn("w:rFonts"))
for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
    rfonts.set(qn(attr), FONT)
szcs = rpr.makeelement(qn("w:szCs"), {qn("w:val"): str(int(SIZE * 2))})
rpr.append(szcs)
style.paragraph_format.space_after = Pt(0)
style.paragraph_format.space_before = Pt(0)
style.paragraph_format.line_spacing = 1.0


def add_runs(p, text, bold=False):
    # **bold** and `code` inline markup
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        b = bold
        if part.startswith("**"):
            part, b = part[2:-2], True
        elif part.startswith("`"):
            part = part[1:-1]
        r = p.add_run(part)
        r.bold = b
        r.font.cs_bold = b


lines = (HERE / "review.md").read_text(encoding="utf-8").splitlines()
i = 0
while i < len(lines):
    line = lines[i].rstrip()
    if not line:
        i += 1
        continue
    if line.startswith("# "):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, line[2:], bold=True)
        p.runs[0].font.size = Pt(SIZE + 2)
    elif line.startswith("## "):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        add_runs(p, line[3:], bold=True)
    elif line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].startswith("|"):
            cells = [c.strip() for c in lines[i].strip("|").split("|")]
            if not set("".join(cells)) <= set("-: "):
                rows.append(cells)
            i += 1
        t = doc.add_table(rows=len(rows), cols=len(rows[0]))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r, row in enumerate(rows):
            for c, val in enumerate(row):
                cp = t.cell(r, c).paragraphs[0]
                if c >= 2:
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_runs(cp, val, bold=(r == 0))
        continue
    elif line.startswith("- "):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Cm(0.6)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs(p, line[2:])
    else:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if line.startswith("**6610"):
            p.paragraph_format.space_before = Pt(3)
        add_runs(p, line)
    i += 1

for p in doc.paragraphs:
    for r in p.runs:
        r.font.name = FONT
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), FONT)

out = HERE / "review.docx"
doc.save(out)
print(out)
